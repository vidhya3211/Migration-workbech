"""Deterministic workbook reader: formulas, labelled constants and defined names, each with a cell locator.

A cell's *label* is the nearest text cell to its left on the same row, else the header above it. Formula
references are rewritten with labels so ``=B2*B3`` reads as ``SUM_ASSURED * BASE_RATE``.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook

_REF = re.compile(r"(?<![A-Za-z_!])(\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z0-9_(])")


def load_workbook_bytes(data: bytes) -> Workbook:
    return load_workbook(io.BytesIO(data), data_only=False, read_only=False)


def cell_repr(cell: Cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, float):
        return format(Decimal(repr(v)).normalize(), "f")
    return str(v)


@dataclass
class SheetFinding:
    kind: str           # calculation | constant
    sheet: str
    cell: str
    label: str
    expression: str
    raw: str
    inputs: list[str]
    confidence: float
    reasons: list[str]


def _label_for(ws, cell: Cell) -> str | None:
    for col in range(cell.column - 1, 0, -1):
        v = ws.cell(row=cell.row, column=col).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    for row in range(cell.row - 1, 0, -1):
        v = ws.cell(row=row, column=cell.column).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    return None


def extract(data: bytes) -> tuple[list[SheetFinding], dict]:
    wb = load_workbook_bytes(data)
    findings: list[SheetFinding] = []
    labels: dict[tuple[str, str], str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or isinstance(cell.value, str) and not cell.value.startswith("="):
                    continue
                label = _label_for(ws, cell)
                if label:
                    labels[(ws.title, cell.coordinate)] = label
    defined = {}
    for name, dn in wb.defined_names.items():
        for sheet, ref in dn.destinations:
            defined[(sheet, ref.replace("$", ""))] = name
            labels.setdefault((sheet, ref.replace("$", "")), name)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                key = (ws.title, cell.coordinate)
                label = labels.get(key)
                if isinstance(v, str) and v.startswith("="):
                    inputs: list[str] = []

                    def sub(m: re.Match[str]) -> str:
                        ref = m.group(1).replace("$", "")
                        name = labels.get((ws.title, ref))
                        if name:
                            inputs.append(name)
                            return name
                        inputs.append(ref)
                        return ref
                    expr = _REF.sub(sub, v[1:])
                    findings.append(SheetFinding(
                        "calculation", ws.title, cell.coordinate, label or f"{ws.title}!{cell.coordinate}", expr, v,
                        list(dict.fromkeys(inputs)), 0.85 if label else 0.6,
                        ["workbook formula"] + ([] if label else ["no label found for the result cell"]),
                    ))
                elif isinstance(v, (int, float)) and not isinstance(v, bool) and label:
                    findings.append(SheetFinding(
                        "constant", ws.title, cell.coordinate, label, cell_repr(cell), cell_repr(cell), [],
                        0.8 if key in defined else 0.7,
                        ["defined name" if key in defined else "labelled numeric input cell"],
                    ))
    summary = {"sheets": wb.sheetnames, "defined_names": sorted({n for n in defined.values()}),
               "formulas": sum(1 for f in findings if f.kind == "calculation"),
               "constants": sum(1 for f in findings if f.kind == "constant")}
    return findings, summary
