"""Regenerates the synthetic data extracts and the rates workbook in the sample estate.

    python samples/build_samples.py

All people and policies are fictitious. Source values follow the legacy (COBOL) premium logic; target
values are what the target platform's own logic produced after RLG's pipelines loaded the records - so
reconciliation finds the real, explainable breaks (admin fee 25 vs 30; age loading at exactly 60).
"""
from __future__ import annotations

import csv
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
D = Decimal

# pol, status, smoker, freq, age, term, sum_assured, base_rate, fund, yif, cust, start, last_prem
POLICIES = [
    ("P000000001", "A", "N", "M", 35, 10, "150000.00", "0.9500", "0.00", 3, "C000000001", "20210315", "20260901"),
    ("P000000002", "A", "Y", "M", 42, 20, "250000.00", "1.1200", "0.00", 6, "C000000002", "20190101", "20260901"),
    ("P000000003", "A", "N", "A", 51, 25, "300000.00", "1.3000", "0.00", 2, "C000000003", "20230610", "20260610"),
    ("P000000004", "L", "N", "M", 47, 15, "100000.00", "1.0500", "0.00", 9, "C000000004", "20150720", "20181101"),
    ("P000000005", "A", "Y", "M", 58, 10, "120000.00", "1.4500", "0.00", 4, "C000000005", "20220201", "20260901"),
    ("P000000006", "P", "N", "M", 63, 20, "80000.00", "1.6000", "0.00", 12, "C000000006", "20130905", "20200105"),
    ("P000000007", "A", "N", "M", 60, 10, "200000.00", "1.5500", "0.00", 1, "C000000007", "20250402", "20260901"),
    ("P000000008", "A", "N", "M", 29, 30, "400000.00", "0.8000", "0.00", 7, "C000000008", "20180812", "20260901"),
    ("P000000009", "X", "N", "M", 44, 20, "175000.00", "1.1000", "0.00", 5, "C000000009", "00000000", "20260901"),
    ("P000000010", "A", "Y", "M", 66, 10, "90000.00", "1.9000", "0.00", 2, "C000000010", "20240118", "20260901"),
    ("P000000011", "A", "N", "M", 38, 15, "220000.00", "1.0000", "0.00", 8, "C000000011", "20170330", "20260901"),
    ("P000000012", "A", "N", "M", 45, 20, "160000.00", "1.1500", "0.00", 3, "C000000012", "20230101", "20260901"),
]

CUSTOMERS = [
    ("C000000001", "Alice Harper", "1990-04-12", "LS1 4AP", "QQ123456A"),
    ("C000000002", "Brian Okafor", "1983-11-02", "M1 2AB", "QQ234567B"),
    ("C000000003", "Carla Mendes", "1974-06-30", "B3 3CD", "QQ345678C"),
    ("C000000004", "Daniel Price", "1978-01-19", "CF10 1EF", "QQ456789D"),
    ("C000000005", "Ellen Brooks", "1967-09-09", "EH1 1GH", "QQ567890A"),
    ("C000000006", "Frank Liddell", "1962-02-28", "BS1 5JK", "QQ678901B"),
    ("C000000007", "Grace Whitfield", "1965-05-21", "NE1 7LM", "QQ789012C"),
    ("C000000008", "Harun Patel", "1996-12-01", "L1 8NP", "QQ890123D"),
    ("C000000009", "Isla Moreno", "1981-07-14", "G1 9QR", "QQ901234A"),
    ("C000000010", "John Smith", "1959-03-02", "SW1A 1AA", "QQ012345B"),
    ("C000000011", "Karen Doyle", "1987-10-25", "NG1 2ST", "QQ112233C"),
    ("C000000012", "Jon Smith", "1959-03-02", "SW1A 1AA", ""),
]


def q2(x: Decimal) -> Decimal:
    return x.quantize(D("0.01"), ROUND_HALF_UP)


def term_factor(term: int) -> Decimal:
    return D("1.000") if term <= 10 else D("1.150") if term <= 20 else D("1.300")


def premium(p, *, admin_fee: Decimal, age_inclusive: bool) -> Decimal:
    _, _, smoker, _, age, term, sa, rate, *_ = p
    risk = D("1")
    if smoker == "Y":
        risk *= D("1.5")
    if (age >= 60) if age_inclusive else (age > 60):
        risk = (risk * D("1.25")).quantize(D("0.0001"), ROUND_HALF_UP)
    base = q2((D(sa) / 1000) * D(rate) * risk * term_factor(term))
    return q2(base / 12) + admin_fee


def main() -> None:
    src, tgt = HERE / "source" / "data", HERE / "target" / "data"
    src.mkdir(parents=True, exist_ok=True)
    tgt.mkdir(parents=True, exist_ok=True)
    with (src / "policies.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["POL_NUMBER", "POL_STATUS", "SMOKER_FLAG", "PAY_FREQ", "HOLDER_AGE", "TERM_YEARS", "SUM_ASSURED",
                    "BASE_RATE", "FUND_VALUE", "YEARS_IN_FORCE", "MONTHLY_PREMIUM", "CUST_ID", "START_DT", "LAST_PREM_DT"])
        for p in POLICIES:
            w.writerow([p[0], p[1], p[2], p[3], p[4], p[5], p[6], p[7], p[8], p[9],
                        premium(p, admin_fee=D("25.00"), age_inclusive=True), p[10], p[11], p[12]])
    with (src / "customers.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["CUST_ID", "CUST_NAME", "DOB", "POSTCODE", "NI_NUMBER"])
        w.writerows(CUSTOMERS)
    with (tgt / "lta_policy.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["policy_ref", "status_code", "is_smoker", "payment_frequency", "holder_age", "term_years",
                    "sum_assured", "base_rate", "fund_value", "years_in_force", "monthly_premium", "party_id",
                    "commencement_date"])
        for p in POLICIES:
            if p[0] == "P000000012":      # not loaded by the pipeline - a missing record
                continue
            start = p[11]
            iso = f"{start[:4]}-{start[4:6]}-{start[6:]}" if start != "00000000" else "1900-01-01"
            w.writerow([p[0], p[1], "true" if p[2] == "Y" else "false", p[3], p[4], p[5], p[6], p[7], p[8], p[9],
                        premium(p, admin_fee=D("30.00"), age_inclusive=False), p[10], iso])
    with (tgt / "party.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["party_id", "full_name", "party_birth_dt", "post_code", "ni_number"])
        w.writerows(CUSTOMERS[:11])

    wb = Workbook()
    ws = wb.active
    ws.title = "Premium"
    ws["A1"], ws["B1"] = "Premium worked example", None
    ws["A2"], ws["B2"] = "Sum assured", 150000
    ws["A3"], ws["B3"] = "Base rate", 0.95
    ws["A4"], ws["B4"] = "Monthly admin fee", 25
    ws["A5"], ws["B5"] = "Smoker loading", 1.5
    ws["A6"], ws["B6"] = "Monthly premium", "=ROUND((B2/1000)*B3/12,2)+B4"
    rates = wb.create_sheet("Rates")
    rates["A1"], rates["B1"] = "Term band", "Term factor"
    rates["A2"], rates["B2"] = "1-10 years", 1.0
    rates["A3"], rates["B3"] = "11-20 years", 1.15
    rates["A4"], rates["B4"] = "21+ years", 1.3
    wb.defined_names["ADMIN_FEE"] = DefinedName("ADMIN_FEE", attr_text="Premium!$B$4")
    (HERE / "source" / "spreadsheets").mkdir(parents=True, exist_ok=True)
    wb.save(HERE / "source" / "spreadsheets" / "LTA_Rates.xlsx")
    print("sample estate regenerated")


if __name__ == "__main__":
    main()
